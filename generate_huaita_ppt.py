from pptx import Presentation
from pptx.util import Pt

OUTPUT = "淮塔（淮海战役烈士纪念塔）课堂汇报.pptx"

slides = [
    ("淮塔（淮海战役烈士纪念塔）", ["课堂汇报", "汇报人：________", "日期：________"]),
    ("淮塔是什么", [
        "全称：淮海战役烈士纪念塔",
        "地点：江苏省徐州市",
        "定位：纪念先烈的革命纪念地与城市地标",
        "核心主题：铭记淮海战役历史与英雄精神",
    ]),
    ("历史背景", [
        "淮海战役是解放战争三大战役之一",
        "为缅怀在战役中牺牲的烈士而建设淮塔",
        "淮塔承载徐州与全国共同的红色记忆",
        "见证人民战争与军民团结的历史力量",
    ]),
    ("建筑特色", [
        "整体形态挺拔庄重，纪念性强",
        "塔体、基座、浮雕等构成完整纪念空间",
        "材质与线条强调肃穆、稳重、永恒",
        "视觉上体现“纪念丰碑”象征",
    ]),
    ("纪念意义", [
        "缅怀淮海战役英烈，致敬牺牲奉献",
        "强化“铭记历史、珍爱和平”的公共记忆",
        "弘扬革命精神与家国情怀",
        "激发新时代青年的责任意识",
    ]),
    ("文化价值", [
        "徐州红色文化的重要代表性符号",
        "连接城市历史文脉与集体记忆",
        "是爱国主义与革命传统教育的重要载体",
        "提升公众对历史文化遗产的认同感",
    ]),
    ("实践意义", [
        "适用于思政课、主题团日、研学实践",
        "帮助学生将历史知识转化为价值认同",
        "通过实地参观与讲解增强学习体验",
        "在纪念活动中发挥社会教育功能",
    ]),
    ("总结", [
        "淮塔不仅是纪念建筑，更是精神丰碑",
        "它记录历史、凝聚记忆、传承信念",
        "我们应铭记先烈、珍惜和平、勇担责任",
    ]),
    ("结束", ["感谢聆听", "欢迎交流" ]),
]

prs = Presentation()

# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = slides[0][0]
slide.placeholders[1].text = "\n".join(slides[0][1])

for title, bullets in slides[1:]:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        for run in p.runs:
            run.font.size = Pt(24)

for slide in prs.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "微软雅黑"

prs.save(OUTPUT)
print(f"Generated: {OUTPUT}")
